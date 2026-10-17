# -*- coding: utf-8 -*-
"""ai_toolbox.py - 补齐 Agent 缺失的基础能力（纯标准库 + 已装库）
用法:
  python ai_toolbox.py fetch <url> <outfile>          抓任意大小网页(突破8KB截断)
  python ai_toolbox.py read  <file> [--head N]        读文本(UTF-8/GBK自动识别)
  python ai_toolbox.py write <file> <content|->       写文本(强制UTF-8, content=- 表示读stdin)
  python ai_toolbox.py grep  <dir> <pattern> [--ext .go,.py]   等价 rg
  python ai_toolbox.py doc   <file>                   解析 pdf/docx/xlsx
"""
import sys, os, re

def _w(s):
    sys.stdout.buffer.write((str(s) + "\n").encode("utf-8", "replace"))

def cmd_fetch(a):
    import urllib.request
    url, out = a[0], a[1]
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 ai_toolbox"})
    with urllib.request.urlopen(req, timeout=30) as r:
        data = r.read(); code = r.status; ct = r.headers.get("Content-Type", "")
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    open(out, "wb").write(data)
    _w("status=%s bytes=%d type=%s -> %s" % (code, len(data), ct, out))

def _txt(p):
    raw = open(p, "rb").read()
    for enc in ("utf-8-sig", "utf-8", "gbk", "latin-1"):
        try: return raw.decode(enc), enc
        except UnicodeDecodeError: pass
    return raw.decode("utf-8", "replace"), "utf-8/replace"

def cmd_read(a):
    p = a[0]; n = int(a[a.index("--head")+1]) if "--head" in a else 0
    text, enc = _txt(p); lines = text.splitlines()
    _w("### %s (enc=%s, lines=%d)" % (p, enc, len(lines)))
    _w("\n".join(lines[:n] if n else lines))

def cmd_write(a):
    p, c = a[0], a[1]
    if c == "-": c = sys.stdin.read()
    d = os.path.dirname(os.path.abspath(p))
    os.makedirs(d, exist_ok=True)
    open(p, "w", encoding="utf-8", newline="\n").write(c)
    _w("written %d chars -> %s" % (len(c), p))

def cmd_grep(a):
    root, pat = a[0], a[1]
    exts = tuple(e.strip() for e in a[a.index("--ext")+1].split(",")) if "--ext" in a else None
    rx = re.compile(pat, re.I); hits = 0
    skip = {".git", "node_modules", "data", "data-m3", "dist", "__pycache__"}
    for dp, dn, fn in os.walk(root):
        dn[:] = [d for d in dn if d not in skip]
        for f in fn:
            if exts and not f.endswith(exts): continue
            p = os.path.join(dp, f)
            try:
                if os.path.getsize(p) > 4*1024*1024: continue
                t, _ = _txt(p)
            except Exception: continue
            for i, ln in enumerate(t.splitlines(), 1):
                if rx.search(ln):
                    _w("%s:%d: %s" % (p, i, ln.strip()[:200])); hits += 1
                    if hits >= 80: _w("... 超80条已截断"); return
    if not hits: _w("no match")

def cmd_doc(a):
    p = a[0]; low = p.lower()
    if low.endswith(".pdf"):
        from PyPDF2 import PdfReader
        r = PdfReader(p); _w("pages=%d" % len(r.pages))
        for i, pg in enumerate(r.pages[:50], 1):
            _w("--- p%d ---\n%s" % (i, (pg.extract_text() or "").strip()))
    elif low.endswith(".docx"):
        import docx
        d = docx.Document(p)
        _w("\n".join(x.text for x in d.paragraphs))
        for ti, tb in enumerate(d.tables, 1):
            _w("--- table%d ---" % ti)
            for row in tb.rows: _w(" | ".join(c.text.strip() for c in row.cells))
    elif low.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        for ws in wb.worksheets:
            _w("--- sheet: %s (%dx%d) ---" % (ws.title, ws.max_row, ws.max_column))
            for row in ws.iter_rows(max_row=200, values_only=True):
                _w(" | ".join("" if v is None else str(v) for v in row))
    else:
        _w("unsupported: %s" % p)

CMDS = {"fetch": cmd_fetch, "read": cmd_read, "write": cmd_write, "grep": cmd_grep, "doc": cmd_doc}
if __name__ == "__main__":
    if len(sys.argv) < 2 or sys.argv[1] not in CMDS:
        _w(__doc__); sys.exit(1)
    CMDS[sys.argv[1]](sys.argv[2:])
